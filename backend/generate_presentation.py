#!/usr/bin/env python3
"""
Generate a Minimalist, Academic, Professional PowerPoint Presentation (.pptx)
for SecureVault: Secure File Sharing System Using Hill Cipher and DES.

Design Principles:
- 16:9 Widescreen format.
- Minimalist palette: Slate Charcoal, Crisp White, Muted Gray, Subtle Navy/Graphite.
- Clean typography and structured layout.
- Detailed technical and mathematical depth suitable for B.Tech Viva defense.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Minimalist Monochrome / Executive Slate Palette
COLOR_BG = RGBColor(248, 249, 250)         # Crisp Minimalist Off-White (#F8F9FA)
COLOR_CARD_BG = RGBColor(255, 255, 255)    # Pure White Card
COLOR_BORDER = RGBColor(226, 232, 240)     # Subtle Border (#E2E8F0)
COLOR_PRIMARY = RGBColor(15, 23, 42)       # Deep Slate (#0F172A)
COLOR_SECONDARY = RGBColor(71, 85, 105)    # Muted Slate (#475569)
COLOR_MUTED = RGBColor(148, 163, 184)      # Light Slate (#94A3B8)
COLOR_ACCENT = RGBColor(30, 41, 59)        # Dark Charcoal/Navy (#1E293B)
COLOR_HIGHLIGHT = RGBColor(37, 99, 235)    # Restrained Blue Accent (#2563EB)
COLOR_ACCENT_BG = RGBColor(241, 245, 249)  # Light Slate Pill/Box

FONT_HEADING = "Helvetica"
FONT_BODY = "Arial"


def create_deck():
    prs = Presentation()
    # 16:9 Widescreen Dimensions (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # completely blank

    def add_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG
        return bg

    def add_header(slide, title_text, category="CRYPTOGRAPHY & NETWORK SECURITY"):
        # Category Tracker Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = FONT_HEADING
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_HIGHLIGHT

        # Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_HEADING
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

        # Thin divider line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.015)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.color.rgb = COLOR_BORDER

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Minimalist & Impactful)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Main Card
    add_card(s1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))

    tbox = s1.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(3.2))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "ACADEMIC B.TECH MINI-PROJECT PRESENTATION"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_HIGHLIGHT
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Secure File Sharing System\nUsing Hill Cipher and DES"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "A Two-Stage Cryptographic Pipeline Operating in Ring Z_256 with DES-CBC and HMAC-SHA256 Authenticated Integrity"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_SECONDARY

    # Metadata / Author footer on card
    meta_box = s1.shapes.add_textbox(Inches(1.8), Inches(4.8), Inches(9.733), Inches(1.2))
    mtf = meta_box.text_frame
    mp = mtf.paragraphs[0]
    mp.text = "Department of Computer Science & Engineering  •  Course: Cryptography & Network Security (CNS)\nProduct Name: SecureVault  •  Demonstration & Viva Defense"
    mp.font.name = FONT_BODY
    mp.font.size = Pt(11)
    mp.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 2: Project Objectives & Problem Definition
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Project Objectives & Problem Statement", "Introduction & Motivation")

    # Column 1: The Problem
    add_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    b1 = s2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "The Problem: Single Cipher Limitations"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    points1 = [
        ("Vulnerability of Single-Layer Ciphers:", "A single symmetric cipher is vulnerable to single-point algorithmic weakness or key leakage."),
        ("Hill Cipher Limitation:", "Classical Hill Cipher is purely linear and easily broken via Known-Plaintext Attacks (KPA) using linear algebra."),
        ("DES Limitation:", "DES possesses a 56-bit effective key length (2^56 keys), which is vulnerable to brute-force attacks via modern hardware."),
        ("Binary File Incompatibility:", "Classical Hill Cipher operates mod 26 (A-Z only) and cannot process binary files (PDFs, images, archives) without severe corruption."),
    ]
    for title, desc in points1:
        pt = t1.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(3)
        pd = t1.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    # Column 2: The Solution
    add_card(s2, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    b2 = s2.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "Proposed Solution: SecureVault Architecture"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    points2 = [
        ("Layered Cryptographic Defense:", "Compounds polygraphic linear matrix diffusion (Hill) with non-linear Feistel block permutation (DES)."),
        ("Ring Z_256 Matrix Extension:", "Formulates Hill Cipher modulo 256, allowing arbitrary 8-bit bytes to be encrypted and decrypted with 100% byte fidelity."),
        ("DES in CBC Mode:", "Utilizes Cipher Block Chaining with random 8-byte IVs to eliminate identical ciphertext blocks."),
        ("PBKDF2 & HMAC Integrity:", "Derives distinct keys using 100,000 rounds of HMAC-SHA256 and seals output via Encrypt-then-MAC (EtM)."),
    ]
    for title, desc in points2:
        pt = t2.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(3)
        pd = t2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 3: End-to-End Cryptographic Architecture
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Two-Stage Cryptographic Pipeline Architecture", "System Design")

    # 4 Architecture Flow Cards
    stages = [
        ("STAGE 01", "Key Derivation", "PBKDF2-HMAC-SHA256", "• 100,000 computation rounds\n• 16-byte random salt\n• Derives 64 pseudorandom bytes:\n  - 8B DES Key\n  - 32B HMAC Key\n  - 4B Hill Matrix Seed"),
        ("STAGE 02", "Layer 1: Hill Cipher", "Linear Matrix in Z_256", "• Byte-pair block grouping\n• Formula: C1 = (P · K^T) mod 256\n• Invertible 2x2 matrix K in Z_256\n• Guaranteed odd determinant\n• Rapid algebraic byte diffusion"),
        ("STAGE 03", "Layer 2: DES Block Cipher", "CBC Mode + PKCS#7", "• 64-bit block size (8 bytes)\n• 16-round Feistel network\n• Random 8-byte IV per file\n• Eliminates structural patterns\n• Formula: C2_i = DES(C1_i ⊕ C2_{i-1})"),
        ("STAGE 04", "Integrity & Packaging", "HMAC-SHA256 (EtM)", "• Encrypt-then-MAC paradigm\n• Sealed in .svault container\n• Constant-time verification\n• Detects tampering / wrong key\n• Prevents padding oracle attacks"),
    ]

    for idx, (st_num, st_name, st_sub, st_desc) in enumerate(stages):
        left = Inches(0.8 + idx * 2.98)
        add_card(s3, left, Inches(1.8), Inches(2.8), Inches(4.9))

        tb = s3.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(2.4), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = st_num
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_HIGHLIGHT
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = st_name
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = st_sub
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_SECONDARY
        p2.space_after = Pt(12)

        p3 = tf.add_paragraph()
        p3.text = st_desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = COLOR_SECONDARY

    # =========================================================================
    # SLIDE 4: Mathematical Deep-Dive: Hill Cipher in Ring Z_256
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "Mathematical Formulation: Hill Cipher in Ring Z_256", "Layer 1 Mathematics")

    add_card(s4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    b1 = s4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "Invertibility Condition Modulo 256"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    math_text = [
        ("Theorem:", "A square matrix K in Z_256 has a modular inverse K^-1 mod 256 if and only if gcd(det(K) mod 256, 256) = 1."),
        ("Prime Factorization:", "Since 256 = 2^8, its sole prime factor is 2. Therefore, gcd(det(K), 256) = 1 strictly holds if and only if det(K) is ODD."),
        ("Matrix Inversion Formula:", "K^-1 = (det(K)^-1 * adj(K)) mod 256\nwhere det(K)^-1 is the modular multiplicative inverse of det(K) in Z_256, and adj(K) is the adjugate matrix."),
        ("Adjugate for 2x2 Matrix:", "For K = [[a, b], [c, d]], adj(K) = [[d, -b], [-c, a]]."),
    ]
    for title, desc in math_text:
        pt = t1.add_paragraph()
        pt.text = f"•  {title}"
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t1.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    add_card(s4, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    b2 = s4.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "Guaranteed Invertible Matrix Derivation"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    proof_steps = [
        ("Parity Enforcement Proof:", "Let 4 seed bytes be transformed as:\n• a = byte[0] | 1       (odd: a = 2k + 1)\n• b = byte[1] & 0xFE    (even: b = 2m)\n• c = byte[2] & 0xFE    (even: c = 2n)\n• d = byte[3] | 1       (odd: d = 2p + 1)"),
        ("Determinant Parity:", "det(K) = (a*d - b*c) mod 256\ndet(K) mod 2 = (odd * odd) - (even * even)\ndet(K) mod 2 = 1 - 0 = 1 (Always Odd!)"),
        ("Conclusion:", "Every matrix generated via this theorem is unconditionally invertible in Z_256 without exceptions."),
        ("Vectorized Execution:", "NumPy vectorization ensures C = (P * K^T) mod 256 operates with O(N) throughput without numeric truncation."),
    ]
    for title, desc in proof_steps:
        pt = t2.add_paragraph()
        pt.text = f"•  {title}"
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Stage 2: DES Block Cipher in CBC Mode
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "Stage 2: Data Encryption Standard (DES) in CBC Mode", "Block Cipher Mechanics")

    add_card(s5, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    b1 = s5.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "DES Specifications & Feistel Network"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    des_points = [
        ("Block Size:", "64 bits (8 bytes) per operational block."),
        ("Key Size:", "64 bits total: 56 effective bits + 8 parity bits."),
        ("Feistel Rounds:", "16 iterative rounds of substitution and permutation."),
        ("Round Operations:", "L_i = R_{i-1}\nR_i = L_{i-1} ⊕ F(R_{i-1}, K_i)\nwhere F uses 8 non-linear S-boxes (48 to 32 bits)."),
        ("PKCS#7 Padding:", "Appends N bytes of value N (1 <= N <= 8) so input length is a strict multiple of 8 bytes."),
    ]
    for title, desc in des_points:
        pt = t1.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t1.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    add_card(s5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    b2 = s5.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "Why CBC Mode (Cipher Block Chaining)?"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    cbc_points = [
        ("ECB Insecurity:", "Electronic Codebook (ECB) encrypts each block independently. Identical plaintext blocks produce identical ciphertexts, exposing file patterns."),
        ("CBC Chaining Formula:", "C_0 = IV (8 random bytes)\nC_i = E_K(P_i ⊕ C_{i-1}) for i >= 1\nEvery block depends on all preceding blocks."),
        ("CBC Decryption Formula:", "P_i = D_K(C_i) ⊕ C_{i-1}\nDecryption is fully parallelizable."),
        ("IV Randomization:", "A unique 8-byte IV is generated per operation. Encrypting the same file twice produces completely distinct ciphertexts."),
    ]
    for title, desc in cbc_points:
        pt = t2.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6: Critical Academic Advisory: Obsoleteness of DES
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Academic Security Advisory: The Obsoleteness of DES", "Cryptanalysis & Ethics")

    add_card(s6, Inches(1.2), Inches(1.7), Inches(10.933), Inches(5.2))
    tb = s6.shapes.add_textbox(Inches(1.6), Inches(2.0), Inches(10.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Why DES Cannot Be Claimed as Modern Production Security"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(14)

    advisory_points = [
        ("Exhaustive Keyspace Vulnerability:", "DES has an effective key length of 56 bits (2^56 ≈ 7.2 × 10^16 possible keys). This keyspace was considered formidable in 1977, but is trivially small today."),
        ("Historical Precedent (EFF DES Cracker, 1998):", "The Electronic Frontier Foundation built 'Deep Crack' for under $250,000, recovering DES keys via brute-force in less than 56 hours."),
        ("Modern Distributed GPU Capabilities:", "With modern GPU clusters (e.g. Hashcat, custom FPGAs), the complete 56-bit DES keyspace can be exhausted in less than a single day."),
        ("Role in SecureVault Mini-Project:", "DES is utilized in this project strictly as an educational vehicle to demonstrate the Feistel structure and CBC block chaining alongside Hill Cipher. It is NOT advertised as modern production security."),
        ("Modern Recommendation:", "Modern enterprise implementations must transition to AES-256 (Advanced Encryption Standard in GCM mode) or ChaCha20-Poly1305."),
    ]
    for title, desc in advisory_points:
        pt = tf.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = tf.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(10)

    # =========================================================================
    # SLIDE 7: Key Management & HMAC Integrity
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Key Derivation (PBKDF2) & Authenticated Integrity (HMAC)", "Defense-in-Depth")

    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    b1 = s7.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "PBKDF2 Key Derivation (NIST SP 800-132)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    pbkdf_points = [
        ("Password Entropy Strengthening:", "Human passwords have low entropy. PBKDF2 maps arbitrary strings to uniform cryptographic keys."),
        ("100,000 Rounds of HMAC-SHA256:", "Applies 100,000 iterations to dramatically slow down GPU-based dictionary and rainbow table attacks."),
        ("16-Byte Cryptographic Salt:", "Generated via os.urandom. Prevents cross-file precomputation attacks."),
        ("64-Byte Derived Stream:", "Outputs partitioned keys:\n• Bytes 00..07: DES Key (8B)\n• Bytes 08..39: HMAC Key (32B)\n• Bytes 40..43: Hill Matrix Seed (4B)"),
    ]
    for title, desc in pbkdf_points:
        pt = t1.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t1.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    add_card(s7, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    b2 = s7.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "HMAC-SHA256 Encrypt-then-MAC (EtM)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    hmac_points = [
        ("Encrypt-then-MAC Paradigm:", "DES-CBC ensures confidentiality, but does NOT guarantee integrity. SecureVault computes an HMAC over the entire container payload."),
        ("Constant-Time Verification:", "Uses hmac.compare_digest to prevent side-channel timing attacks during authentication."),
        ("Padding Oracle Defense:", "If HMAC signature check fails, decryption aborts immediately without calling the DES unpadding routine."),
        ("Unified Error Contract:", "Returns 'Decryption failed: invalid key or corrupted file' to prevent information leakage."),
    ]
    for title, desc in hmac_points:
        pt = t2.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = t2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: The .svault Binary Container Format
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Binary Container Specification (.svault Format)", "Data Structures")

    add_card(s8, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.133), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Documented Binary Container Header & Payload Structure"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    fields = [
        ("Magic Identifier", "6 Bytes", "ASCII string b'SVAULT' to identify file type unambiguously"),
        ("Version & Flags", "2 Bytes", "Version 0x01 + 0x00 reserved flags byte"),
        ("PBKDF2 Salt", "16 Bytes", "Cryptographically secure random salt for key derivation"),
        ("DES Initialization Vector", "8 Bytes", "Random 8-byte IV for CBC block chaining"),
        ("Hill Cipher Coefficients", "5 Bytes", "Dimension (2) + coefficients a, b, c, d (uint8)"),
        ("Original File Size", "8 Bytes", "uint64 big-endian representing exact unpadded file length"),
        ("Filename Descriptor", "2B + Variable", "Length (uint16) + UTF-8 original filename bytes"),
        ("DES Ciphertext", "8B + Variable", "Ciphertext length (uint64) + PKCS#7 padded ciphertext bytes"),
        ("HMAC-SHA256 Integrity Tag", "32 Bytes", "Cryptographic signature computed over all preceding bytes"),
    ]

    for name, sz, desc in fields:
        pt = tf.add_paragraph()
        pt.text = f"•  {name} ({sz}): "
        pt.font.bold = True
        pt.font.size = Pt(10.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(1)
        pd = tf.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(4)

    # =========================================================================
    # SLIDE 9: Experimental Verification & Test Suite
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Experimental Verification & Quality Assurance", "Validation & Testing")

    add_card(s9, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    b1 = s9.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "18 Automated Pytest Test Cases (100% Pass)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    tests = [
        ("Hill Cipher Mod 256:", "Passes arbitrary byte values (0..255) with exact modular inversion."),
        ("DES-CBC Cipher:", "Correct PKCS#7 block alignment and unpadding."),
        ("Text File (.txt):", "Byte-for-byte fidelity confirmed via SHA-256."),
        ("Document (.pdf):", "PDF binary headers and EOF markers restored 1:1."),
        ("Image (.png/.jpg):", "PNG magic bytes and pixel buffers preserved."),
        ("Archive (.zip):", "ZIP archives decompress cleanly after roundtrip."),
        ("Edge Cases:", "0-byte empty file and 1 MB large binary file pass."),
    ]
    for title, desc in tests:
        pt = t1.add_paragraph()
        pt.text = f"[✓]  {title} "
        pt.font.bold = True
        pt.font.size = Pt(10.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(1)
        pd = t1.add_paragraph()
        pd.text = f"     {desc}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(4)

    add_card(s9, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    b2 = s9.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "Tamper Detection & Live Demo Results"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    demo_pts = [
        ("Wrong Password Rejection:", "Throws 400 Bad Request with immediate HMAC failure."),
        ("Ciphertext Bit-Flipping:", "Single-bit alteration detected and rejected without executing DES unpadding."),
        ("Corrupted Container:", "Truncated file rejected by container length validator."),
        ("Live Terminal Demo (demo_test.py):", "Reads demo.txt, encrypts, decrypts, and confirms:\n  Original SHA256  == Decrypted SHA256\n  Result: [ PASS ]"),
    ]
    for title, desc in demo_pts:
        pt = t2.add_paragraph()
        pt.text = f"[✓]  {title} "
        pt.font.bold = True
        pt.font.size = Pt(10.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(1)
        pd = t2.add_paragraph()
        pd.text = f"     {desc}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(5)

    # =========================================================================
    # SLIDE 10: Viva Defense Preparation (Examiner FAQ)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)
    add_header(s10, "Viva Examiner Defense: Key Questions & Answers", "Viva Preparation")

    add_card(s10, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2))
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.133), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Core Theoretical Questions Addressed by SecureVault"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    qa = [
        ("Q1: Why can classical Hill Cipher NOT be used directly on binary files?",
         "Classical Hill Cipher uses mod 26 for letters A-Z. Binary files contain 256 distinct byte values (0..255). Expanding to Ring Z_256 ensures zero information loss."),
        ("Q2: What is the exact mathematical condition for a matrix to be invertible in Z_256?",
         "gcd(det(K) mod 256, 256) = 1. Since 256 = 2^8, the determinant det(K) MUST BE ODD."),
        ("Q3: Why is CBC mode preferred over ECB mode in DES?",
         "In ECB mode, identical plaintext blocks produce identical ciphertext blocks (e.g. ECB penguin). In CBC mode, each block is XORed with the previous ciphertext block."),
        ("Q4: Why is HMAC-SHA256 required if DES already encrypts the file?",
         "Encryption provides confidentiality, but does NOT guarantee integrity. HMAC guards against ciphertext bit-flipping and padding oracle attacks in constant time."),
    ]
    for q, a in qa:
        pt = tf.add_paragraph()
        pt.text = f"•  {q}"
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = tf.add_paragraph()
        pd.text = f"   {a}"
        pd.font.size = Pt(10)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(10)

    # =========================================================================
    # SLIDE 11: Conclusion & Future Scope
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_background(s11)
    add_header(s11, "Conclusion & Future Research Directions", "Summary")

    add_card(s11, Inches(1.2), Inches(1.7), Inches(10.933), Inches(5.2))
    tb = s11.shapes.add_textbox(Inches(1.6), Inches(2.0), Inches(10.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Summary of Contributions"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    c_points = [
        ("Successful Pedagogical Hybrid Cryptosystem:", "Demonstrated that combining algebraic matrix substitution (Hill Cipher in Z_256) with Feistel block permutations (DES-CBC) provides layered security principles."),
        ("Guaranteed Reversibility & Integrity:", "All operations guarantee 100% byte-for-byte fidelity with cryptographic SHA-256 audit verification."),
        ("Production-Grade Engineering:", "Clean monorepo architecture with Next.js 16, FastAPI, SQLite, zero external cloud dependencies, and full test suite."),
        ("Future Scope — Post-Quantum & Modern Ciphers:", "Replace DES with AES-256-GCM or ChaCha20-Poly1305 for production applications, and investigate Lattice-based Post-Quantum Cryptography (PQC) matrix replacements for Hill Cipher."),
    ]
    for title, desc in c_points:
        pt = tf.add_paragraph()
        pt.text = f"•  {title} "
        pt.font.bold = True
        pt.font.size = Pt(11.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.space_after = Pt(2)
        pd = tf.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_SECONDARY
        pd.space_after = Pt(10)

    output_path = os.path.join(os.path.dirname(__file__), "..", "SecureVault_Presentation.pptx")
    output_path = os.path.abspath(output_path)
    prs.save(output_path)
    print(f"Successfully generated presentation: {output_path}")
    return output_path


if __name__ == "__main__":
    create_deck()
